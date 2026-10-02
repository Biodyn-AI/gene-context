"""Additional file 1 (Excel): the full results behind the multi-model comparison, one sheet per analysis, read from
results/*.json. Run with a Python that has openpyxl. Out: bmc_submission/Additional_file_1.xlsx"""
import os, json, glob
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE); RES = os.path.join(os.path.dirname(ROOT), "results")
MODELS = [("MaxToki-217M", "ctx217m600", [4, 2]), ("MaxToki-1B", "ctx1b", [7, 4]), ("scGPT", "ctx_scgpt600", [4, 2, 9]),
          ("STATE (SE-600M)", "ctx_state600", [6, 3, 12]), ("Expression only, landmarks", "ctxexpr", [0]),
          ("Expression only, principal components", "ctxexprpc", [0]), ("Expression only, centroids", "ctxexprcen50c", [0]),
          ("Expression only, centroids with all expressing cells (no 50-cell cap; check only)", "ctxexprcen50", [0]),
          ("Expression only, centroids with 512 PCs and no cap (fails the positive control; not used)", "ctxexprcen", [0]),
          ("MaxToki-217M, untrained", "ctxrand", [4]), ("scGPT, untrained", "ctx_scgpt600rand", [4]),
          ("STATE, untrained", "ctx_state600rand", [6])]
AXN = {"nuclear_vs_surface": "Nuclear / surface", "mito_vs_cytoskeleton": "Mitochondrion / cytoskeleton",
       "transcription_vs_transport": "Transcription / transport", "cellcycle_vs_diff": "Cell cycle / differentiation"}


def J(name):
    p = os.path.join(RES, name)
    return json.load(open(p)) if os.path.exists(p) else None


def r3(x):
    return None if x is None else round(float(x), 4)


def rp(x):                       # p values: 3 significant digits, never rounded to 0
    return None if x is None else float(f"{float(x):.3g}")


wb = Workbook(); ws0 = wb.active; ws0.title = "README"
ws0.append(["Additional file 1. Full results of the multi-model comparison (all representations on the same 600 cells per "
            "cell type, partitions, 5000-gene panel and occurrence cap 50)."])
ws0.append(["Each sheet is generated from the result files in results/ of the gene-context repository by "
            "bmc_submission/build/make_additional_file.py."])
ws0.append(["EXCESS = mean cosine between a gene's relative context shift in the two cell partitions, minus the same for a "
            "randomly paired gene (Methods). Main-effect replication = cosine between the two partitions' shared shift of all "
            "genes, the positive control that the measurement works."])
ws0.append(["min. cell types = genes count-balanced in at least this many of the 12 cell types; covariate = the gene "
            "abundance covariate used in the rank regression (MaxToki mean rank on the same cells, or mean "
            "log(1 + counts per 10,000))."])
ws0.append(["Not every analysis was run on every representation: the untrained models went through EXCESS, the curated-target "
            "test and (untrained scGPT and STATE only) the functional axes; the rank regression of EXCESS was run at each trained "
            "model's depth-matched layer only; empty cells mean the analysis was not run."])
ws0.append(["Functional-axis, co-expression, tightness and directional-congruence p values are empirical one-sided "
            "(k + 1)/(N + 1) over N = 300 null axes. Curated-target p values and the steering 'Sign p (random)' are one-sided "
            "sign tests over TFs or cells. The steering fixed-split p is (k + 1)/(N + 1) over N = 30 random splits (two-sided, "
            "on the size of the swing). Cross-model functional-z is against random-gene axes (number given in the Methods)."])
ws0.append(["The steering and prediction-link sheets also include the 1000-cell MaxToki-217M runs (files without a "
            "ctx217m600 or ctx1b tag)."])

ws = wb.create_sheet("EXCESS and rank control")
ws.append(["Representation", "Layer", "EXCESS", "Same gene", "Different gene", "Cell-type pairs", "Main-effect replication",
           "EXCESS after rank regression", "% surviving"])
for lab, pre, taps in MODELS:
    pol = J(f"ctx_polysemy__{pre}.json"); pc = J(f"ctx_position_confound__{pre}.json")
    for t in taps:
        L = f"L{t:02d}"
        if not pol or L not in pol["taps"]: continue
        r = pol["taps"][L]; rr = (pc or {}).get("taps", {}).get(L, {})
        ws.append([lab, t, r3(r["excess"]), r3(r["same"]), r3(r["diff"]), r["n_pairs"], r3(r.get("main_effect_replication")),
                   r3(rr.get("excess_residualised")), round(100 * rr["excess_residualised"] / r["excess"]) if rr else None])

ws = wb.create_sheet("Functional axes and nulls")
ws.append(["Representation", "Layer", "Min. cell types", "Covariate", "Axis", "Validity AUC", "Functional-z",
           "Functional-z, rank-controlled", "Power", "p, size+co-expression matched", "p, strong co-expression modules",
           "p, size+tightness matched"])
for lab, pre, taps in MODELS:
    for mc in ("6", "9"):
        for cov, sfx in (("MaxToki mean rank", ""), ("mean log(1 + counts per 10,000)", "__cov-expr")):
            tag = f"__{pre}__minctx{mc}{sfx}"
            fa = J(f"ctx_functional_axes{tag}.json"); v3 = J(f"ctx_coexpr_null_v3{tag}.json")
            ti = J(f"ctx_tightness_null{tag}.json")
            if not fa: continue
            for t in taps:
                L = f"L{t:02d}"
                if L not in fa["taps"]: continue
                for a, r in fa["taps"][L].items():
                    s3 = (v3 or {}).get("summary", {}).get(a, {}) if (v3 and t == taps[0]) else {}
                    a3 = (v3 or {}).get("axes", {}).get(a, {}) if (v3 and t == taps[0]) else {}
                    tt = (ti or {}).get("taps", {}).get(L, {}).get("axes", {}).get(a, {})
                    ws.append([lab, t, int(mc), cov, AXN.get(a, a), r3(r["validity_auc"]), r3(r["z_raw"]),
                               r3(r["z_rank_controlled"]), r3(a3.get("power")), rp(s3.get("p_matched")) if s3 else None,
                               rp(s3.get("p_strong_modules_v2")) if s3 else None, rp(tt.get("matched_p")) if tt else None])

ws = wb.create_sheet("Context-specific function")
ws.append(["Representation", "Layer", "Test", "Axis", "Statistic", "Value", "p", "Null mean", "Null SD"])
for lab, pre, taps in MODELS:
    dp = J(f"ctx_directional_probe__{pre}__minctx6.json")
    if dp:
        for L, tv in dp["taps"].items():
            for a, r in tv["axes"].items():
                ws.append([lab, int(L[1:]), "Directional congruence", AXN.get(a, a), "beta", r3(r["beta"]), rp(r["p"]),
                           r3(r["null_mean"]), r3(r["null_sd"])])
    ct = J(f"ctx_curated_targets__{pre}.json")
    if ct:
        for L, r in ct["taps"].items():
            ws.append([lab, int(L[1:]), f"Curated targets, static closeness ({r['n_tfs']} TFs)", "", "mean cosine excess",
                       r3(r["static_excess_mean"]), rp(r["static_sign_p"]), None, None])
            ws.append([lab, int(L[1:]), f"Curated targets, modulation ({r['n_modul']} TFs)", "", "mean Spearman rho",
                       r3(r["modulation_mean_rho"]), rp(r["modulation_sign_p"]), None, None])
            if r.get("specific_n"):
                ws.append([lab, int(L[1:]), f"Curated targets, own vs other TFs' targets ({r['specific_n_pos']} of "
                           f"{r['specific_n']} TFs positive)", "", "mean cosine excess", r3(r["specific_excess_mean"]),
                           rp(r["specific_sign_p"]), None, None])

ws = wb.create_sheet("Cross-model")
ws.append(["Comparison", "Representation", "Layer", "Cell types", "Genes count-balanced in at least one cell type", "Width", "Cell-type pairs", "EXCESS",
           "EXCESS 95% CI low", "EXCESS 95% CI high", "Main-effect replication", "Functional-z (nuclear/surface)"])
for tag, name in (("all600", "own count-balanced genes"), ("all600_common", "genes shared by all representations")):
    cm = J(f"ctx_cross_model__{tag}.json")
    if not cm: continue
    for lab, r in cm.items():
        if not isinstance(r, dict) or "excess" not in r: continue
        fz = (r["func_z"].get("nuclear_vs_surface") or {}).get("z")
        ws.append([name, lab, r.get("tap"), r["n_ctx"], r.get("n_genes_balanced", r["n_genes"]), r["dim"], r.get("pairs_scored"), r3(r["excess"]),
                   r3(r["excess_ci"][0]), r3(r["excess_ci"][1]), r3(r["main_effect_replication"]), r3(fz)])

ws = wb.create_sheet("Steering")
ws.append(["File", "Model", "Axis", "Layer", "Push unit", "Strength (x residual norm)", "Functional swing", "SEM",
           "Random-direction swing", "Cells func > random", "Sign p (random)", "Fixed random splits",
           "Mean |split swing|", "Splits with |swing| >= functional", "p (fixed splits)", "z (fixed splits)"])
for p in sorted(glob.glob(os.path.join(RES, "ctx_causal*.json"))):
    d = json.load(open(p)); f = os.path.basename(p)
    ax = "nuclear/surface" if "mito" not in f and "trans" not in f else ("mitochondrion/cytoskeleton" if "mito" in f else "transcription/transport")
    sn = d.get("split_null") or {}
    for a, m in d["signed"].items():
        on = sn and float(a.split("_")[1]) == sn.get("alpha")
        ws.append([f, d.get("model", "217m"), ax, d["site"] + 1, d.get("alpha_unit", "raw"), float(a.split("_")[1]),
                   r3(m["func_swing"]), r3(m["func_swing_sem"]), r3(m["rand_swing"]), m["func_gt_rand"], rp(m["sign_p"]),
                   sn.get("n_splits") if on else None, r3(sn.get("split_abs_mean")) if on else None,
                   sn.get("n_splits_abs_ge_func") if on else None, rp(sn.get("p_abs")) if on else None,
                   r3(sn.get("z_abs")) if on else None])

ws = wb.create_sheet("Prediction link")
ws.append(["File", "Model", "Contextualisation from", "Donor", "Genes", "Comparisons", "Mean benefit (nats)",
           "Median benefit (nats)", "Partial Spearman", "95% CI low", "95% CI high"])
for p in sorted(glob.glob(os.path.join(RES, "ctx_prediction_link*.json"))):
    d = json.load(open(p)); f = os.path.basename(p)
    if "arms" in d:
        for arm, r in d["arms"].items():
            if "mean_benefit" not in r: continue
            ws.append([f, d["model"], d["extraction"], arm.replace("_", " "), r["n_genes"], r["n_comparisons"],
                       r3(r["mean_benefit"]), r3(r["median_benefit"]), r3(r["partial_rho"]), r3(r["partial_ci"][0]),
                       r3(r["partial_ci"][1])])
    else:
        ws.append([f, d.get("model", "217m"), d.get("extraction", "ctx_maxtoki_L04"), "random donor (original design)",
                   d["n_genes"], d.get("n_comparisons"), r3(d["mean_context_benefit"]), r3(d.get("median_context_benefit")),
                   r3(d["partial_rho"]), r3(d["partial_ci"][0]), r3(d["partial_ci"][1])])

for w in wb.worksheets:
    for c in w[1]:
        c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical="top")
    for col in w.columns:
        w.column_dimensions[col[0].column_letter].width = 18
out = os.path.join(ROOT, "Additional_file_1.xlsx"); wb.save(out)
print("wrote", out, [(w.title, w.max_row - 1) for w in wb.worksheets])
