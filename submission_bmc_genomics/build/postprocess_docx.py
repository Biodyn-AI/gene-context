"""Post-process a pandoc DOCX for BMC: continuous line numbers, page numbers in the footer, single-spaced
tables with plain black borders and no shading. Usage: postprocess_docx.py in.docx out.docx"""
import sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING

d = Document(sys.argv[1])
for sec in d.sections:
    sp = sec._sectPr
    for old in sp.findall(qn("w:lnNumType")):
        sp.remove(old)
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1"); ln.set(qn("w:restart"), "continuous"); ln.set(qn("w:distance"), "283")
    # lnNumType must come before pgNumType/cols/docGrid etc.; insert after pgMar if present
    pgmar = sp.find(qn("w:pgMar"))
    (pgmar.addnext(ln) if pgmar is not None else sp.insert(0, ln))
    footer = sec.footer; footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.text = ""; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for tag, txt in [("begin", None), (None, "PAGE"), ("end", None)]:
        if tag:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), tag); run._r.append(fc)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; run._r.append(it)
    run.font.name = "Times New Roman"; run.font.size = Pt(12)

def set_borders(tbl):
    tblPr = tbl._tbl.tblPr
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ["top", "bottom", "insideH"]:
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:color"), "000000")
        b.append(e)
    tblPr.append(b)
for tbl in d.tables:
    set_borders(tbl)
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            for sh in tcPr.findall(qn("w:shd")):
                tcPr.remove(sh)
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for r in par.runs:
                    r.font.size = Pt(10)
# no hard page breaks anywhere (BMC: "do not use page breaks")
n_breaks = 0
for br in d.element.body.iter(qn("w:br")):
    if br.get(qn("w:type")) == "page":
        br.getparent().remove(br); n_breaks += 1
d.save(sys.argv[2])
print(f"wrote {sys.argv[2]}: {len(d.sections)} section(s), {len(d.tables)} table(s), removed {n_breaks} page break(s)")
